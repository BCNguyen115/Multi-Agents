"""Document Loader & Section Analyzer for legal documents.

Recursively traverses the dataset directory, extracts text from PDF and DOCX
files, and analyzes document structure to identify section boundaries
(ARTICLE, SECTION, numbered headings, etc.) for downstream section-based
chunking.

Usage:
    from src.ingestion.document_loader import load_and_analyze_documents
    documents = load_and_analyze_documents("/app/dataset")
"""

from __future__ import annotations

import hashlib
import logging
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from pypdf import PdfReader

from src.ingestion.ocr import ocr_pdf_pages
from src.shared.dataset_reader import decode_bytes
from src.shared.logger import get_logger

logger: logging.Logger = get_logger(__name__)

# ---------------------------------------------------------------------------
# Data classes
# ---------------------------------------------------------------------------


@dataclass
class Section:
    """A single identified section within a document.

    Attributes:
        title: The heading / identifier of the section (e.g. "Article 1").
        content: The full text content of the section.
        start_index: Character offset where this section starts in the raw text.
        end_index: Character offset where this section ends in the raw text.
    """

    title: str
    content: str
    start_index: int
    end_index: int


@dataclass
class SectionedDocument:
    """A document that has been loaded and analyzed for section boundaries.

    Attributes:
        filename: Original filename (e.g. ``contract.pdf``).
        filepath: Absolute path to the source file.
        category: Subdirectory name (e.g. ``msa``, ``nda``).
        raw_text: The complete extracted text.
        sections: Ordered list of identified sections.
        detected_pattern: Name of the regex pattern family that matched best.
        metadata: Arbitrary key-value metadata carried through the pipeline.
    """

    filename: str
    filepath: str
    category: str
    raw_text: str
    sections: list[Section] = field(default_factory=list)
    detected_pattern: str = "none"
    metadata: dict[str, Any] = field(default_factory=dict)


# ---------------------------------------------------------------------------
# Section heading patterns (ordered by specificity / priority)
# ---------------------------------------------------------------------------

SECTION_PATTERNS: list[tuple[re.Pattern[str], str]] = [
    # ARTICLE I / ARTICLE 1 / ARTICLE 10
    (
        re.compile(
            r"(?m)^[ \t]*(ARTICLE\s+(?:[IVXLCDM]+|\d+)[\.:\s\u2014\u2013-]*.*)$",
            re.IGNORECASE,
        ),
        "ARTICLE",
    ),
    # SECTION 1.01 / Section 1 / SECTION 2
    (
        re.compile(
            r"(?m)^[ \t]*((?:SECTION|Sec\.?)\s+\d+(?:\.\d+)*[\.:\s\u2014\u2013-]*.*)$",
            re.IGNORECASE,
        ),
        "SECTION",
    ),
    # Numbered heading: 1. DEFINITIONS / 2. Term
    (
        re.compile(
            r"(?m)^[ \t]*(\d{1,2}\.[ \t]+[A-Z].{2,80})$",
        ),
        "NUMBERED_HEADING",
    ),
    # Decimal sub-sections: 1.1 / 2.3.1
    (
        re.compile(
            r"(?m)^[ \t]*(\d{1,2}\.\d{1,2}(?:\.\d{1,2})?[ \t]+.{2,80})$",
        ),
        "DECIMAL_SUBSECTION",
    ),
    # ALL-CAPS headings (at least 6 chars, on their own line)
    (
        re.compile(
            r"(?m)^[ \t]*([A-Z][A-Z \t]{5,})$",
        ),
        "ALL_CAPS_HEADING",
    ),
]


# Formats that are not contracts get their own heading patterns (used only for them: PDF/DOCX detection is unchanged).
MARKDOWN_PATTERNS: list[tuple[re.Pattern[str], str]] = [
    (re.compile(r"(?m)^[ \t]{0,3}(#{1,6}[ \t]+\S.{0,200})$"), "MARKDOWN_HEADING"),
    *SECTION_PATTERNS,  # a plain .txt (or a .md without #) may still be numbered like a contract
]
SLIDE_PATTERNS: list[tuple[re.Pattern[str], str]] = [(re.compile(r"(?m)^(Slide \d+:.*)$"), "SLIDE")]

PRIMARY_EXTENSIONS: tuple[str, ...] = (".pdf", ".docx")
EXTRA_EXTENSIONS: tuple[str, ...] = (".pptx", ".txt", ".md")  # inside category folders only (never loose files like a report)
SUPPORTED_EXTENSIONS: tuple[str, ...] = PRIMARY_EXTENSIONS + EXTRA_EXTENSIONS


# ---------------------------------------------------------------------------
# Text extraction helpers
# ---------------------------------------------------------------------------


_REPEATED_LINE_SHARE: float = 0.5   # a line on more than this share of pages is a running header/footer
_MIN_PAGES_FOR_STRIPPING: int = 4
_MAX_HEADER_LINE_CHARS: int = 120


_PAGE_NUMBER_LINE = re.compile(r"^(?:page\s*)?\d+(?:\s*(?:of|/)\s*\d+)?$|^page\b", re.IGNORECASE)


def _line_key(line: str) -> str:
    """Comparison key of a line. Only page-number lines (``3``, ``Page 3 of 10``) get their digits masked, so
    real content that merely has the same wording on many pages is not mistaken for a running header."""
    text: str = line.strip().lower()
    return re.sub(r"\d+", "#", text) if _PAGE_NUMBER_LINE.match(text) else text


def _strip_repeated_lines(pages: list[str]) -> list[str]:
    """Drop running headers/footers/page numbers (short lines present on most pages).

    They would otherwise be read as section headings (ALL-CAPS pattern) and pollute every chunk.
    """
    if len(pages) < _MIN_PAGES_FOR_STRIPPING:
        return pages
    presence: dict[str, int] = {}
    for page in pages:
        for key in {_line_key(l) for l in page.splitlines() if 3 <= len(l.strip()) <= _MAX_HEADER_LINE_CHARS}:
            presence[key] = presence.get(key, 0) + 1
    limit: float = max(3.0, _REPEATED_LINE_SHARE * len(pages))
    running: set[str] = {key for key, count in presence.items() if count > limit}
    if not running:
        return pages
    return ["\n".join(l for l in page.splitlines() if _line_key(l) not in running) for page in pages]


def _join_pages(pages: list[str]) -> tuple[str, list[int]]:
    """Join pages with newlines; ``starts[i]`` is the character offset where page ``i + 1`` begins."""
    starts: list[int] = []
    offset: int = 0
    parts: list[str] = []
    for page in pages:
        starts.append(offset)
        parts.append(page)
        offset += len(page) + 1
    return "\n".join(parts), starts


_MIN_TEXT_CHARS_PER_PAGE: int = 20  # below this on average a PDF has no usable text layer: it is a scan


def _extract_pdf(filepath: str) -> tuple[str, list[int], bool]:
    """Extract text from a PDF with pypdf; a scan (no text layer) is read by OCR when that is available.

    Returns:
        ``(text, page_starts, ocr_used)``; empty text on failure or when a scan cannot be read.
    """
    try:
        reader: PdfReader = PdfReader(filepath)
        pages: list[str] = [(page.extract_text() or "").strip() for page in reader.pages]
    except Exception as exc:
        logger.warning("Failed to read PDF %s: %s", filepath, exc, extra={"session_id": "INGESTION"})
        return "", [], False
    ocr_used = False
    if pages and sum(len(p) for p in pages) / len(pages) < _MIN_TEXT_CHARS_PER_PAGE:
        try:
            ocr_pages = ocr_pdf_pages(filepath)
        except Exception as exc:  # noqa: BLE001 - a failing OCR is the same as no OCR: the caller reports "no text"
            logger.warning("OCR of %s failed: %s", filepath, exc, extra={"session_id": "INGESTION"})
            ocr_pages = []
        if sum(len(p) for p in ocr_pages) > sum(len(p) for p in pages):
            pages, ocr_used = ocr_pages, True
    text, starts = _join_pages(_strip_repeated_lines(pages))
    return text, starts, ocr_used


def _extract_pdf_text(filepath: str) -> tuple[str, list[int]]:
    """``(text, page_starts)`` of a PDF (see ``_extract_pdf``); empty text on failure."""
    text, starts, _ = _extract_pdf(filepath)
    return text, starts


def _extract_plain_text(filepath: str) -> str:
    """A ``.txt`` / ``.md`` file, in whatever encoding it was saved (BOMs, UTF-16, Windows code pages)."""
    try:
        return decode_bytes(Path(filepath).read_bytes())[0].replace("\r\n", "\n").replace("\r", "\n")
    except OSError as exc:
        logger.warning("Failed to read %s: %s", filepath, exc, extra={"session_id": "INGESTION"})
        return ""


def _extract_pptx_text(filepath: str) -> str:
    """A PowerPoint deck as text: one ``Slide N: <title>`` block per slide (text boxes, tables, speaker notes)."""
    try:
        from pptx import Presentation

        deck = Presentation(filepath)
    except Exception as exc:  # noqa: BLE001 - corrupt file, or python-pptx missing
        logger.warning("Failed to read PPTX %s: %s", filepath, exc, extra={"session_id": "INGESTION"})
        return ""
    blocks: list[str] = []
    for number, slide in enumerate(deck.slides, start=1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip().replace("\n", " ")
        lines: list[str] = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                lines += [p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()]
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Notes: {notes}")
        if title or lines:
            blocks.append(f"Slide {number}: {title}".rstrip() + "\n" + "\n".join(lines))
    return "\n".join(blocks)


def _extract_docx_text(filepath: str) -> str:
    """Extract all text from a DOCX file using python-docx.

    Args:
        filepath: Absolute or relative path to the DOCX.

    Returns:
        Concatenated paragraph text, or empty string on failure.
    """
    try:
        doc: DocxDocument = DocxDocument(filepath)
        paragraphs: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(paragraphs)
    except Exception as exc:
        logger.warning(
            "Failed to read DOCX %s: %s",
            filepath,
            exc,
            extra={"session_id": "INGESTION"},
        )
        return ""


# ---------------------------------------------------------------------------
# Section analysis
# ---------------------------------------------------------------------------


def _detect_best_pattern(
    text: str, patterns: list[tuple[re.Pattern[str], str]] | None = None
) -> tuple[re.Pattern[str] | None, str]:
    """Determine which section-heading pattern best fits the document.

    Strategy: try each pattern in priority order and pick the one with the
    most matches (minimum 2 to avoid false positives).

    Args:
        text: The full document text.

    Returns:
        ``(compiled_pattern, pattern_name)``, or ``(None, "none")`` when no pattern matches at least twice
        (the chunker then uses its fixed-size fallback).
    """
    best_pattern: re.Pattern[str] | None = None
    best_name: str = "none"
    best_count: int = 0

    for pattern, name in patterns or SECTION_PATTERNS:
        matches: list[str] = pattern.findall(text)
        if len(matches) >= 2 and len(matches) > best_count:
            best_count = len(matches)
            best_pattern = pattern
            best_name = name

    return best_pattern, best_name


def _split_into_sections(
    text: str,
    pattern: re.Pattern[str] | None,
    pattern_name: str,
) -> list[Section]:
    """Split document text into sections using the detected heading pattern.

    Args:
        text: Full document text.
        pattern: The compiled regex for section headings (``None``: no structure detected).
        pattern_name: Human-readable name for logging.

    Returns:
        Ordered list of ``Section`` objects.
    """
    matches: list[re.Match[str]] = list(pattern.finditer(text)) if pattern is not None else []

    if not matches:
        return [
            Section(
                title="Full Document",
                content=text.strip(),
                start_index=0,
                end_index=len(text),
            )
        ]

    sections: list[Section] = []

    # Capture preamble text before first heading
    if matches[0].start() > 0:
        preamble_text: str = text[: matches[0].start()].strip()
        if preamble_text:
            sections.append(
                Section(
                    title="Preamble",
                    content=preamble_text,
                    start_index=0,
                    end_index=matches[0].start(),
                )
            )

    # Split on each heading
    for i, match in enumerate(matches):
        heading: str = match.group(1).strip()
        start: int = match.start()
        end: int = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        section_text: str = text[start:end].strip()

        sections.append(
            Section(
                title=heading,
                content=section_text,
                start_index=start,
                end_index=end,
            )
        )

    return sections


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def load_document(filepath: str, category: str, filename: str | None = None) -> SectionedDocument | None:
    """Extract text, detect the section structure and build the identity metadata of ONE document.

    The folder ingestion and the chat upload both go through here, so a document gets the same ``doc_key``,
    ``doc_hash``, page map and sections whichever way it enters the knowledge base. ``None`` when the extension is
    not supported or no text could be extracted (e.g. a scanned PDF).
    """
    filename = filename or os.path.basename(filepath)
    ext: str = os.path.splitext(filename)[1].lower()
    logger.info("Loading: [%s] %s", category, filename, extra={"session_id": "INGESTION"})

    # Extract text (+ page offsets for PDFs, so chunks can cite a page)
    page_starts: list[int] = []
    ocr_used: bool = False
    patterns: list[tuple[re.Pattern[str], str]] | None = None
    raw_text: str
    if ext == ".pdf":
        raw_text, page_starts, ocr_used = _extract_pdf(filepath)
    elif ext == ".docx":
        raw_text = _extract_docx_text(filepath)
    elif ext == ".pptx":
        raw_text, patterns = _extract_pptx_text(filepath), SLIDE_PATTERNS
    elif ext in (".txt", ".md"):
        raw_text, patterns = _extract_plain_text(filepath), MARKDOWN_PATTERNS
    else:
        return None

    if not raw_text.strip():
        logger.warning("Empty text extracted from %s — skipping", filename, extra={"session_id": "INGESTION"})
        return None

    # Analyze section structure
    pattern, pattern_name = _detect_best_pattern(raw_text, patterns)
    sections: list[Section] = _split_into_sections(raw_text, pattern, pattern_name)

    return SectionedDocument(
        filename=filename,
        filepath=filepath,
        category=category,
        raw_text=raw_text,
        sections=sections,
        detected_pattern=pattern_name,
        metadata={
            "source": filepath,
            "filename": filename,
            "category": category,
            "doc_key": f"{category}/{filename}",
            "doc_hash": hashlib.sha256(raw_text.encode("utf-8")).hexdigest(),
            "page_starts": page_starts,
            "total_chars": len(raw_text),
            "num_sections": len(sections),
            "ocr": ocr_used,
        },
    )


def load_and_analyze_documents(
    dataset_dir: str,
    supported_extensions: tuple[str, ...] = SUPPORTED_EXTENSIONS,
) -> list[SectionedDocument]:
    """Load all documents from a dataset directory and analyze their structure.

    Recursively walks ``dataset_dir``, extracts text from each supported file,
    detects the best section-heading pattern, and splits the text into
    ordered sections.

    Args:
        dataset_dir: Root directory containing categorised document folders.
        supported_extensions: File extensions to process.

    Returns:
        List of ``SectionedDocument`` objects ready for chunking.

    Raises:
        FileNotFoundError: If ``dataset_dir`` does not exist.
    """
    dataset_path: Path = Path(dataset_dir).resolve()
    if not dataset_path.is_dir():
        raise FileNotFoundError(f"Dataset directory not found: {dataset_path}")

    documents: list[SectionedDocument] = []
    skipped: int = 0

    for root, _dirs, files in os.walk(dataset_path):
        for filename in sorted(files):
            ext: str = os.path.splitext(filename)[1].lower()
            if ext not in supported_extensions:
                skipped += 1
                continue

            category: str = Path(root).name
            if category == dataset_path.name:
                category = "root"
                if ext in EXTRA_EXTENSIONS:  # a loose .md/.txt next to the folders (e.g. the ingestion report) is not a document
                    skipped += 1
                    continue

            doc: SectionedDocument | None = load_document(os.path.join(root, filename), category)
            if doc is None:
                skipped += 1
                continue
            documents.append(doc)

    logger.info(
        "Loaded %d documents (%d skipped) from %s",
        len(documents),
        skipped,
        dataset_path,
        extra={"session_id": "INGESTION"},
    )
    return documents
