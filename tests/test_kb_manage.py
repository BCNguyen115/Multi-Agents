"""Knowledge base: more formats, OCR, listing/deleting documents, page preview, gateway messages in two languages."""
import asyncio
import io
import re
import shutil
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
from fastapi import FastAPI
from PIL import Image, ImageDraw

from src.config import settings
from src.gateway import knowledge
from src.ingestion import document_loader, ocr
from src.ingestion.document_loader import SUPPORTED_EXTENSIONS, load_and_analyze_documents, load_document
from src.ingestion.upload import MAX_SECTION, UploadError, build_upload_chunks, ingest_upload, safe_filename
from src.shared import messages
from src.shared.messages import EN, VI, msg, pick_lang, reset_lang, set_lang

from tests.test_kb_upload import FakeEmbedder, prose


# ------------------------------------------------------------------ more formats
def test_the_supported_formats_are_the_documented_ones():
    assert SUPPORTED_EXTENSIONS == (".pdf", ".docx", ".pptx", ".txt", ".md")


def test_markdown_is_split_at_its_headings_and_keeps_vietnamese(tmp_path):
    body = "\n\n".join(f"# Điều khoản {i}\n{prose(260)}" for i in range(1, 4))
    path = tmp_path / "huong-dan.md"
    path.write_text("﻿" + body, encoding="utf-8")  # with a BOM, as Notepad saves it
    doc = load_document(str(path), "policy")
    assert doc.detected_pattern == "MARKDOWN_HEADING" and len(doc.sections) == 3
    assert doc.sections[0].title.startswith("# Điều khoản 1") and "﻿" not in doc.raw_text
    chunks, report = build_upload_chunks(doc)
    assert report.sections == 3 and all(len(c.raw_content) <= MAX_SECTION for c in chunks)


def test_a_text_file_in_an_old_windows_code_page_is_read(tmp_path):
    path = tmp_path / "note.txt"
    path.write_bytes(("Điều 1. " + prose(120) + "\nĐiều 2. " + prose(120)).encode("cp1258", errors="replace"))
    assert load_document(str(path), "policy") is not None


def make_pptx(path: Path) -> None:
    from pptx import Presentation

    deck = Presentation()
    for number in (1, 2):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"Kế hoạch quý {number}"
        slide.placeholders[1].text_frame.text = prose(260)
        slide.notes_slide.notes_text_frame.text = f"Ghi chú diễn giả {number}"
    deck.save(str(path))


def test_a_powerpoint_becomes_one_section_per_slide_with_notes(tmp_path):
    path = tmp_path / "deck.pptx"
    make_pptx(path)
    doc = load_document(str(path), "plans")
    assert doc.detected_pattern == "SLIDE" and [s.title for s in doc.sections] == ["Slide 1: Kế hoạch quý 1", "Slide 2: Kế hoạch quý 2"]
    assert "Ghi chú diễn giả 2" in doc.raw_text


def test_the_folder_ingestion_reads_the_new_formats_inside_category_folders_but_not_loose_files(tmp_path):
    (tmp_path / "policy").mkdir()
    (tmp_path / "policy" / "a.md").write_text("# A\n" + prose(300) + "\n# B\n" + prose(300), encoding="utf-8")
    (tmp_path / "walkthrough_rag.md").write_text("# report\n" + prose(300) + "\n# more\n" + prose(300), encoding="utf-8")  # the ingestion's own report
    found = {(d.category, d.filename) for d in load_and_analyze_documents(str(tmp_path))}
    assert found == {("policy", "a.md")}


def test_upload_accepts_the_new_extensions_and_names_them_in_the_error():
    for name in ("x.pptx", "x.md", "x.TXT"):
        assert safe_filename(name) == name
    with pytest.raises(UploadError, match="PPTX"):
        safe_filename("x.exe")


def test_a_markdown_upload_is_stored_like_any_other_document(tmp_path):
    data = ("# One\n" + prose(300) + "\n# Two\n" + prose(300)).encode()
    embedder = FakeEmbedder()
    result = asyncio.run(ingest_upload(embedder, "guide.md", data, category="policy", dataset_dir=str(tmp_path)))
    assert result["status"] == "added" and result["chunks"] == 2 and (tmp_path / "policy" / "guide.md").exists()


# ------------------------------------------------------------------ OCR
def scanned_pdf(path: Path, text: str = "SERVICE AGREEMENT TERM 24 MONTHS") -> None:
    """A PDF that is only a picture of text (no text layer), like a scan."""
    image = Image.new("RGB", (1240, 400), "white")
    ImageDraw.Draw(image).text((40, 150), text, fill="black", font_size=48)
    image.save(path, "PDF", resolution=150)


def test_a_scan_is_read_by_ocr_when_the_text_layer_is_empty(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    scanned_pdf(path)
    monkeypatch.setattr(document_loader, "ocr_pdf_pages", lambda p: ["ARTICLE 1 Term\n" + prose(300) + "\nARTICLE 2 Fees\n" + prose(300)])
    doc = load_document(str(path), "scans")
    assert doc is not None and doc.metadata["ocr"] is True and doc.detected_pattern == "ARTICLE"


def test_a_scan_without_ocr_is_refused_as_before(tmp_path, monkeypatch):
    path = tmp_path / "scan.pdf"
    scanned_pdf(path)
    monkeypatch.setattr(document_loader, "ocr_pdf_pages", lambda p: [])
    assert load_document(str(path), "scans") is None
    with pytest.raises(UploadError):
        asyncio.run(ingest_upload(FakeEmbedder(), "scan.pdf", path.read_bytes(), category="scans"))


def test_a_pdf_with_a_text_layer_never_goes_to_ocr(monkeypatch):
    sample = Path("dataset/corporate/community_ex1001.pdf")
    if not sample.exists():
        pytest.skip("sample PDF not in this checkout")
    monkeypatch.setattr(document_loader, "ocr_pdf_pages", lambda p: pytest.fail("OCR must not run on a PDF that has text"))
    doc = load_document(str(sample), "corporate")
    assert doc is not None and doc.metadata["ocr"] is False


@pytest.mark.skipif(not ocr.ocr_available(), reason="needs the tesseract program (installed in the backend image)")
def test_tesseract_really_reads_a_scanned_page(tmp_path):
    path = tmp_path / "scan.pdf"
    scanned_pdf(path)
    text = " ".join(ocr.ocr_pdf_pages(str(path), langs="eng")).upper()
    assert "AGREEMENT" in text and "24" in text


def test_ocr_reports_itself_unavailable_when_switched_off(monkeypatch):
    monkeypatch.setattr(settings, "OCR_ENABLED", False)
    assert ocr.ocr_available() is False and ocr.ocr_pdf_pages("x.pdf") == []


# ------------------------------------------------------------------ gateway messages
def test_both_languages_have_the_same_messages_with_the_same_placeholders():
    assert set(VI) == set(EN)
    slots = lambda text: sorted(re.findall(r"\{\w+\}", text))  # noqa: E731
    assert all(slots(VI[k]) == slots(EN[k]) for k in VI), [k for k in VI if slots(VI[k]) != slots(EN[k])]


def test_the_language_comes_from_the_header_and_defaults_to_vietnamese():
    assert [pick_lang(h) for h in ("en", "en-US,en;q=0.9", "vi-VN", "fr", "", None)] == ["en", "en", "vi", "vi", "vi", "vi"]
    assert msg("login.bad") == VI["login.bad"]
    token = set_lang("en")
    try:
        assert msg("rate.limited", limit=5, seconds=7) == "You are sending requests too fast (at most 5 per minute). Please try again in 7 seconds."
    finally:
        reset_lang(token)
    assert messages.current_lang() == "vi"


# ------------------------------------------------------------------ list / delete / page
class FakePg:
    def __init__(self):
        self.rows = [
            {"doc_key": "nda/a.pdf", "filename": "a.pdf", "category": "nda", "chunks": 5, "pages": 3, "chars": 900, "updated_at": None},
            {"doc_key": "msa/b.docx", "filename": "b.docx", "category": "msa", "chunks": 2, "pages": None, "chars": 400, "updated_at": None},
        ]
        self.deleted = []

    async def fetch(self, sql, *args, **kw):
        return [r for r in self.rows if args[0] in (None, r["category"])]

    async def execute(self, sql, *args, **kw):
        self.deleted.append(args[0])
        hits = [r for r in self.rows if r["doc_key"] == args[0]]
        return f"DELETE {sum(r['chunks'] for r in hits)}"


def client(monkeypatch, tmp_path, roles=None):
    from src.shared import auth

    monkeypatch.setattr(settings, "KNOWLEDGE_DIR", str(tmp_path))
    app = FastAPI()
    app.include_router(knowledge.router)
    app.state.pg_client = FakePg()
    app.state.knowledge_store = SimpleNamespace(invalidate_categories=lambda: None)
    principal = auth.Principal("amy", "t", "d", frozenset(roles or []), authenticated=roles is not None)
    app.dependency_overrides[auth.authenticate] = lambda: principal
    return app, httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://t")


def test_documents_are_listed_with_whether_their_pdf_is_kept(monkeypatch, tmp_path):
    (tmp_path / "nda").mkdir()
    (tmp_path / "nda" / "a.pdf").write_bytes(b"%PDF")

    async def go():
        app, c = client(monkeypatch, tmp_path)
        async with c:
            rows = (await c.get("/api/knowledge/documents")).json()
            assert {r["doc_key"]: r["has_pdf"] for r in rows} == {"nda/a.pdf": True, "msa/b.docx": False}
            assert [r["doc_key"] for r in (await c.get("/api/knowledge/documents", params={"category": "msa"})).json()] == ["msa/b.docx"]

    asyncio.run(go())


def test_deleting_removes_the_chunks_and_the_kept_copy_but_only_for_those_who_may(monkeypatch, tmp_path):
    (tmp_path / "nda").mkdir()
    copy = tmp_path / "nda" / "a.pdf"
    copy.write_bytes(b"%PDF")

    async def go():
        app, c = client(monkeypatch, tmp_path, roles=["analyst"])
        async with c:
            assert (await c.delete("/api/knowledge/documents", params={"doc_key": "nda/a.pdf"})).status_code == 403
            assert copy.exists() and app.state.pg_client.deleted == []
        app, c = client(monkeypatch, tmp_path, roles=["admin"])
        async with c:
            done = (await c.delete("/api/knowledge/documents", params={"doc_key": "nda/a.pdf"})).json()
            assert done == {"doc_key": "nda/a.pdf", "deleted_chunks": 5, "file_removed": True} and not copy.exists()
            assert (await c.delete("/api/knowledge/documents", params={"doc_key": "nda/ghost.pdf"})).status_code == 404

    asyncio.run(go())


@pytest.mark.parametrize("bad", ["../etc/passwd", "nda/../../x.pdf", "nda/a/b.pdf", "x", "nda\\a.pdf", "/abs/a.pdf", "nda/.."])
def test_a_document_key_can_never_point_outside_the_dataset_folder(monkeypatch, tmp_path, bad):
    async def go():
        app, c = client(monkeypatch, tmp_path)
        async with c:
            assert (await c.delete("/api/knowledge/documents", params={"doc_key": bad})).status_code in (404, 422)
            assert (await c.get("/api/knowledge/page", params={"doc_key": bad, "page": 1})).status_code in (404, 422)
        assert app.state.pg_client.deleted == []

    asyncio.run(go())


def test_a_page_comes_back_as_an_image_and_unknown_documents_or_pages_are_404(monkeypatch, tmp_path):
    (tmp_path / "scans").mkdir()
    scanned_pdf(tmp_path / "scans" / "s.pdf")
    (tmp_path / "scans" / "w.docx").write_bytes(b"x")

    async def go():
        app, c = client(monkeypatch, tmp_path)
        async with c:
            ok = await c.get("/api/knowledge/page", params={"doc_key": "scans/s.pdf", "page": 1, "q": "Hop dong dich vu so 42"})
            assert ok.status_code == 200 and ok.headers["content-type"] == "image/png" and ok.content[:4] == b"\x89PNG"
            assert ok.headers["x-page-count"] == "1" and ok.headers["x-highlighted"] == "0"  # a scan has no text layer to search
            assert (await c.get("/api/knowledge/page", params={"doc_key": "scans/s.pdf", "page": 9})).status_code == 404
            assert (await c.get("/api/knowledge/page", params={"doc_key": "scans/w.docx", "page": 1})).status_code == 404  # only PDFs have pages
            assert (await c.get("/api/knowledge/page", params={"doc_key": "scans/none.pdf", "page": 1})).status_code == 404

    asyncio.run(go())


def test_the_cited_passage_is_highlighted_on_a_real_pdf(monkeypatch, tmp_path):
    sample = Path("dataset/corporate/community_ex1001.pdf")
    if not sample.exists():
        pytest.skip("sample PDF not in this checkout")
    (tmp_path / "corporate").mkdir()
    shutil.copy(sample, tmp_path / "corporate" / sample.name)
    snippet = "Section 1.01 Surviving Company . Subject to the terms and provisions of this Agreement and Plan of Merger, and in accordance with the Wyoming Act"

    async def go():
        app, c = client(monkeypatch, tmp_path)
        async with c:
            hit = await c.get("/api/knowledge/page", params={"doc_key": "corporate/" + sample.name, "page": 2, "q": snippet})
            miss = await c.get("/api/knowledge/page", params={"doc_key": "corporate/" + sample.name, "page": 3, "q": snippet})
            assert hit.headers["x-highlighted"] == "1" and miss.headers["x-highlighted"] == "0"  # right page highlighted, wrong page not
            assert Image.open(io.BytesIO(hit.content)).size[0] > 800

    asyncio.run(go())


def test_the_upload_message_follows_the_language_and_mentions_ocr():
    result = {"filename": "a.pdf", "category": "nda", "status": "added", "chunks": 4, "sections": 5, "merged": 1, "split": 1,
              "duplicates_removed": 0, "replaced_chunks": 0, "saved_to_dataset": True, "ocr": True}
    vietnamese = knowledge.describe_upload(result)
    token = set_lang("en")
    try:
        english = knowledge.describe_upload(result)
    finally:
        reset_lang(token)
    assert "thêm mới" in vietnamese and "OCR" in vietnamese
    assert english.startswith("Added **a.pdf**") and "read by OCR" in english and "split into several passages" in english
