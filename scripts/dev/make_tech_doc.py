"""Build docs/Enterprise_Multi_Agent_System_Technical_Documentation.docx (Vietnamese). Run make_diagrams.py first."""
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAME = "Enterprise_Multi_Agent_System"
OUT = Path("docs") / f"{NAME}_Technical_Documentation.docx"
DIAG = Path("temp_diagrams")
NAVY, SLATE, TEAL, GREY = "1B365D", "336699", "008080", "F4F6F9"
FONT = "Calibri"

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(1)
sec.top_margin = sec.bottom_margin = Inches(0.9)


def rgb(h):
    return RGBColor.from_string(h)


def style(name, size, bold=False, color=None, before=0, after=6):
    s = doc.styles[name]
    s.font.name = FONT
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    s.font.size, s.font.bold = Pt(size), bold
    if color:
        s.font.color.rgb = rgb(color)
    s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
    s.paragraph_format.line_spacing = 1.15


style("Normal", 11)
style("Title", 26, True, NAVY)
style("Heading 1", 18, True, NAVY, 18, 8)
style("Heading 2", 14, True, SLATE, 12, 6)
style("Heading 3", 12, True, TEAL, 10, 4)
style("List Bullet", 11)


def shade(cell, color):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(sh)


def borders(cell, color="B8C2D0", sz=4, left=None):
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        c, s = (left, 36) if side == "left" and left else (color, sz)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(s))
        e.set(qn("w:color"), c)
        b.append(e)
    cell._tc.get_or_add_tcPr().append(b)


def margins(cell, m=80):
    mar = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(m))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    cell._tc.get_or_add_tcPr().append(mar)


def h(text, lvl=1):
    doc.add_heading(text, lvl)


def p(text, bold=False, italic=False, align=None, size=None, color=None):
    para = doc.add_paragraph()
    r = para.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = rgb(color)
    if align is not None:
        para.alignment = align
    return para


def bullets(items):
    for it in items:
        para = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            para.add_run(it[0] + ": ").bold = True
            para.add_run(it[1])
        else:
            para.add_run(it)


def table(headers, rows, widths=None, font=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, hd in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(hd)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(font), rgb("FFFFFF")
        shade(c, NAVY)
        borders(c)
        margins(c)
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(str(v)).font.size = Pt(font)
            shade(cells[i], GREY if n % 2 else "FFFFFF")
            borders(cells[i])
            margins(cells[i])
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()


def callout(kind, text):
    color = {"Lưu ý": SLATE, "Cảnh báo": "B26B00", "Best practice": TEAL}[kind]
    c = doc.add_table(rows=1, cols=1).rows[0].cells[0]
    shade(c, GREY)
    borders(c, "FFFFFF", 0, left=color)
    margins(c, 120)
    c.text = ""
    r = c.paragraphs[0].add_run(kind + ": ")
    r.bold, r.font.color.rgb = True, rgb(color)
    c.paragraphs[0].add_run(text)
    doc.add_paragraph()


fig_n = 0


def figure(path, caption, width=6.3):
    global fig_n
    fig_n += 1
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(path), width=Inches(width))
    p(f"Hình {fig_n}: {caption}", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, color="555555")


def code(text):
    c = doc.add_table(rows=1, cols=1).rows[0].cells[0]
    shade(c, GREY)
    borders(c)
    margins(c, 100)
    c.text = ""
    r = c.paragraphs[0].add_run(text)
    r.font.name, r.font.size = "Consolas", Pt(8.5)
    doc.add_paragraph()


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- cover
for _ in range(5):
    doc.add_paragraph()
p("TÀI LIỆU KỸ THUẬT", bold=True, size=12, color=TEAL)
doc.add_paragraph("Enterprise Multi-Agent System", style="Title")
p("Hệ thống đa tác tử doanh nghiệp: vòng lặp Plan – Execute – Verify, RAG lai, phân tích dữ liệu có kiểm chứng và Human-in-the-Loop",
  size=14, color=SLATE)
doc.add_paragraph()
table(["Hạng mục", "Giá trị"], [
    ["Phiên bản tài liệu", "1.0"], ["Tác giả", "BCNguyen115"], ["Ngày", date.today().strftime("%d/%m/%Y")],
    ["Kho mã", "multi_agent_mvp (nhánh main)"],
], [2.0, 4.0], 11)
p("FastAPI  •  LangGraph  •  PostgreSQL + pgvector  •  Redis  •  Next.js  •  LiteLLM  •  Langfuse  •  Docker Compose",
  bold=True, size=10, color=NAVY)
page_break()

h("Mục lục")
for s in ["1. Tóm tắt điều hành", "2. Công nghệ và môi trường", "3. Kiến trúc hệ thống", "4. Module lõi và logic nghiệp vụ",
          "5. Kiến trúc cơ sở dữ liệu và từ điển dữ liệu", "6. Đặc tả API và tích hợp", "7. Giao diện người dùng",
          "8. Bảo mật, hiệu năng và khả năng mở rộng", "9. Triển khai và vận hành", "10. Lộ trình và hạn chế đã biết"]:
    p(s)
page_break()

# ---------------------------------------------------------------- 1
h("1. Tóm tắt điều hành")
h("1.1 Bối cảnh và vấn đề", 2)
p("Doanh nghiệp có tri thức nằm rải rác ở ba nơi: tài liệu (hợp đồng NDA/MSA/SOW, quy trình, slide), dữ liệu bảng (CSV/Excel) và cơ sở "
  "dữ liệu vận hành. Một chatbot LLM đơn lẻ trả lời nhanh nhưng có ba rủi ro: bịa số liệu (hallucination), rò rỉ dữ liệu nhạy cảm giữa các "
  "phòng ban/tenant và thực hiện thao tác ghi ngoài ý muốn. Hệ thống này giải quyết bằng cách tách việc thành nhiều tác tử chuyên biệt, "
  "đặt mọi câu trả lời sau một bước kiểm chứng độc lập và chặn các thao tác nhạy cảm cho tới khi con người phê duyệt.")
h("1.2 Mục tiêu kinh doanh", 2)
bullets([("Độ tin cậy", "đưa tỉ lệ hallucination về gần 0 nhờ kiểm toán số liệu (dashboard) và bắt buộc trích dẫn [n] (RAG)."),
         ("Một cửa vào tri thức", "một khung chat định tuyến tới RAG, phân tích dữ liệu, tìm kiếm web, truy vấn DB và tích hợp hệ thống ngoài."),
         ("Quản trị", "xác thực JWT, phân quyền theo vai trò, Row-Level Security theo tenant/phòng ban, phê duyệt thao tác nhạy cảm."),
         ("Quan sát được", "mọi lượt chạy được truy vết trên Langfuse; có bộ đánh giá offline (LLM-as-a-Judge) và đánh giá truy hồi RAG làm cổng CI.")])
h("1.3 Người dùng mục tiêu", 2)
table(["Nhóm", "Nhu cầu", "Tính năng chính"], [
    ["Nhân viên nghiệp vụ", "Hỏi đáp tài liệu nội bộ, có nguồn", "Chat RAG, xem trang PDF được trích dẫn"],
    ["Chuyên viên dữ liệu", "Tải CSV và nhận dashboard tin cậy", "Data agent, dashboard v2, lọc chéo, SQL trong trình duyệt"],
    ["Quản trị/Approver", "Kiểm soát thao tác nhạy cảm", "HITL approve/reject, quản lý kho tri thức"],
    ["DevOps/IT", "Vận hành, giám sát", "Docker Compose, /ready, Langfuse, CI"],
], [1.5, 2.3, 2.7])
h("1.4 Thành tựu chính", 2)
bullets(["Vòng lặp PEV trên LangGraph với tối đa 2 lần thử lại và cảnh báo khi vẫn không đạt.",
         "Verifier dashboard tái nạp tệp, biên dịch lại từng biểu đồ, tính lại từng KPI và đối chiếu mọi con số trong lời kể: khi đạt thì không cần LLM judge.",
         "RAG lai: vector (HNSW) + full-text (GIN) hợp nhất bằng RRF, rerank cross-encoder có ngân sách thời gian, nạp tài liệu tăng dần và nguyên tử.",
         "Phòng thủ nhiều lớp: quét prompt injection, canary token, SQL allow-list bằng AST, vai trò Postgres chỉ đọc cộng RLS, sandbox cô lập cho mã do LLM viết.",
         "Chạy nhiều bản sao backend sau nginx với trạng thái chia sẻ qua Redis/Postgres."])
page_break()

# ---------------------------------------------------------------- 2
h("2. Công nghệ và môi trường")
table(["Lớp", "Công nghệ", "Lý do chọn"], [
    ["Frontend", "Next.js (App Router), React, Tailwind CSS, ECharts, AG Grid, DuckDB-WASM, react-markdown", "Route handler giữ JWT trong cookie httpOnly; ECharts vẽ dashboard; DuckDB-WASM chạy SQL ngay trên trình duyệt, không gửi dữ liệu đi."],
    ["Backend", "Python 3.11/3.12, FastAPI, Uvicorn, sse-starlette, pydantic-settings", "Async end-to-end, OpenAPI tự sinh, SSE cho luồng PEV thời gian thực."],
    ["Điều phối", "LangGraph (StateGraph), LangChain core", "Máy trạng thái planner→executor→verifier với cạnh điều kiện quay lại."],
    ["LLM", "LiteLLM + OpenRouter; FAST_LLM_MODEL (mặc định openai/gpt-4o-mini), HEAVY_LLM_MODEL (openai/gpt-4o)", "Phân tầng mô hình: nhanh cho planner/verifier, mạnh cho tổng hợp; đổi nhà cung cấp bằng cấu hình."],
    ["CSDL", "PostgreSQL + pgvector (asyncpg, psycopg), Alembic", "Một CSDL cho vector, full-text, hội thoại và bộ nhớ dài hạn; HNSW/GIN; RLS gốc."],
    ["Cache/Trạng thái", "Redis 8", "Lịch sử phiên, bản ghi HITL, giới hạn tốc độ, tệp CSV đang hoạt động; chia sẻ giữa các bản sao."],
    ["Truy hồi", "TEI (HuggingFace text-embeddings-inference) cross-encoder", "Rerank chính xác hơn; có circuit-breaker (RERANKER_TIMEOUT) để không làm chậm."],
    ["Bộ nhớ", "mem0ai (pgvector)", "Bộ nhớ dài hạn theo <tenant>:<user>, có che PII trước khi lưu."],
    ["Dữ liệu", "pandas, numpy, scipy, scikit-learn, pyarrow, openpyxl", "Hồ sơ hóa cột, thống kê, đọc Parquet/Excel."],
    ["Tài liệu", "pypdf, pypdfium2, python-docx, python-pptx, pytesseract (vie+eng)", "Nạp PDF/DOCX/PPTX; OCR cho PDF quét; dựng trang PDF có tô sáng trích dẫn."],
    ["Bảo mật", "PyJWT, cryptography, sqlglot", "Xác minh JWT (HS256/JWKS), kiểm tra SQL theo AST."],
    ["Web", "Tavily, Crawl4AI", "Tìm kiếm web cho search_agent."],
    ["Quan sát", "Langfuse v2 (web + worker)", "Truy vết LLM từng yêu cầu; tự cấp org/project/khóa khi khởi tạo."],
    ["DevOps", "Docker (multi-stage, non-root), Docker Compose, nginx, GitHub Actions, ruff, pytest, vitest, Playwright", "Tái lập môi trường, kiểm thử tự động, cổng chất lượng RAG."],
], [1.1, 2.6, 2.8], 9)
callout("Lưu ý", "requirements.lock ghim mọi phiên bản (đóng băng từ image Linux). Cài bằng: pip install -r requirements-dev.txt -c requirements.lock.")
page_break()

# ---------------------------------------------------------------- 3
h("3. Kiến trúc hệ thống")
p("Hệ thống gồm năm tầng. Client (Next.js) không gọi backend trực tiếp cho các route cần xác thực: mỗi route handler dưới app/api/** đọc JWT từ "
  "cookie httpOnly và chuyển thành header Bearer. Gateway FastAPI xác thực, giới hạn tốc độ, quét prompt injection rồi giao cho Orchestrator. "
  "Orchestrator chạy vòng lặp PEV và gọi tác tử phù hợp; tác tử dùng hạ tầng chung (PostgreSQL, Redis, reranker, sandbox, LLM).")
figure(DIAG / "system_architecture.png", "Kiến trúc tổng thể")
h("3.1 Vai trò thành phần", 2)
table(["Thành phần", "Vị trí", "Vai trò"], [
    ["Gateway", "src/gateway/main.py", "Điểm vào /api/*; xác thực, rate limit, guardrail, chuyển SSE/JSON."],
    ["Orchestrator", "src/orchestrator/core.py (+ approvals.py, streaming.py, common.py)", "Xây StateGraph PEV; định tuyến; HITL; luồng SSE."],
    ["Verifier", "src/orchestrator/verifier.py", "Kiểm toán dashboard; phán quyết RAG; LLM judge cho kết quả khác."],
    ["AgentRegistry", "src/registry/manager.py", "Cổng xác thực khi đăng ký: loại mô tả trùng >80% cosine, thăm dò chức năng."],
    ["Agents", "src/agents/*", "rag, data, search, db, integration; đều kế thừa BaseAgent."],
    ["Ingestion", "src/ingestion/*", "Nạp tài liệu, chia chunk, embedding, OCR, xem trước trang."],
    ["Shared", "src/shared/*", "Auth, security, llm_client, redis/postgres client, rate_limit, tracing, memory."],
], [1.3, 2.4, 2.8])
h("3.2 Vòng lặp PEV", 2)
p("Đồ thị LangGraph có ba nút: planner → executor → verifier; từ verifier có cạnh điều kiện quay lại executor khi kiểm chứng thất bại (tối đa 2 lần), "
  "sau đó trả kết quả kèm biểu ngữ cảnh báo. Không dùng checkpointer: mỗi yêu cầu tự dựng toàn bộ trạng thái, ngữ cảnh hội thoại lấy từ lịch sử Redis.")
table(["Nút", "Nhiệm vụ", "Chi tiết"], [
    ["Planner", "Chọn target_agent và lập plan", "Thứ tự ưu tiên: agent_mode cưỡng bức → có CSV (data_agent) → tách câu hỏi ghép thành DAG song song → LLM planner (FAST_LLM_MODEL) → định tuyến dự phòng theo từ khóa."],
    ["Executor", "Gọi tác tử", "Tra AgentRegistry, chèn phản hồi verifier khi thử lại, chặn HITL trước thao tác nhạy cảm."],
    ["Verifier", "Kiểm chứng đầu ra", "Dashboard: kiểm toán chặt. RAG: phán quyết từ verification{status,grounded,cited}. Khác: LLM judge. Bỏ qua khi từ chối, rỗng, lỗi hoặc đang chờ duyệt."],
], [1.1, 1.8, 3.6])
page_break()

# ---------------------------------------------------------------- 4
h("4. Module lõi và logic nghiệp vụ")
figure(DIAG / "data_flow.png", "Luồng xử lý chính của một lượt chat (SSE)")
figure(DIAG / "component_diagram.png", "Phân rã module và tương tác")
h("4.1 RAG Agent", 2)
p("Nạp tài liệu (scripts.run_ingestion) là tăng dần: mỗi tài liệu có doc_key (<category>/<filename>) và doc_hash; tệp không đổi bị bỏ qua, tệp đổi được thay thế "
  "nguyên tử trong một giao dịch. Số trang PDF được ánh xạ tới chunk; header/footer lặp lại bị loại. Truy hồi gồm: top-k vector (câu hỏi gốc và/hoặc đoạn HyDE) + "
  "top-k full-text, hợp nhất RRF (RAG_RRF_K), suy luận bộ lọc category từ câu hỏi, ngưỡng tương đồng vector (RAG_MIN_VECTOR_SCORE), sau đó rerank TEI theo ngân sách "
  "thời gian tổng.")
bullets([("Planner RAG", "một lần gọi mô hình nhanh biến câu hỏi + lịch sử thành câu hỏi độc lập và đoạn HyDE; lỗi thì dùng câu hỏi nguyên văn."),
         ("Trả lời", "nguồn được bọc trong phong bì có nonce (chỉ loại bỏ tín hiệu injection mạnh); LLM bắt buộc trích dẫn [n]; JSON trả về chứa verification{status,grounded,cited,unsupported_numbers}."),
         ("Thử lại", "NOT_FOUND trong category được nêu tên thì thử lại một lần trên toàn kho."),
         ("Contextual retrieval", "tùy chọn KB_LLM_CONTEXT: một câu ngữ cảnh do LLM sinh được lưu ở đầu content."),
         ("Tải lên từ chat", "POST /api/knowledge/upload (PDF/DOCX/PPTX/TXT/MD); PDF quét đi qua OCR; chunk 200–600 ký tự; bản sao lưu ở KNOWLEDGE_DIR để --prune/--reset không xóa.")])
callout("Best practice", "Các thử nghiệm đã đo và bị loại (dịch truy vấn sang tiếng Anh, rerank pool 20, rerank hai giai đoạn, giới hạn mỗi tệp, mở rộng chunk lân cận, khử trùng lặp liên tài liệu) "
        "không nên đưa lại khi chưa có số liệu mới; xem docs/RAG_REVIEW_AND_FIXES.md và chạy python -m scripts.run_rag_eval.")
h("4.2 Data Agent", 2)
p("Đường ống tất định gồm các mô-đun nhỏ: ingest (CSV/TSV/Excel/Parquet/JSON, mã hóa, tiền tệ/phần trăm/ngày dạng chữ → kiểu chuẩn, mẫu giới hạn) → profiler (vai trò cột theo giá trị/thống kê, "
  "chính sách sum vs average, độ mịn thời gian, chất lượng) → insights (chỉ sự kiện có số) → charts (ChartSpec + compile_chart, nơi duy nhất tính dữ liệu biểu đồ) → story (câu dựng từ sự kiện; "
  "LLM chỉ được diễn đạt lại nếu giữ nguyên tên và số) → verifier. Đầu ra là dashboard_spec v2: lưới 12 cột với story, facts, kpis, charts, slicers, table.")
bullets(["Câu hỏi tiếp theo đi qua qa: LLM viết pandas → sandbox AST allow-list → chạy trong container python-sandbox (HMAC, Parquet thay pickle, mạng nội bộ, fail closed) hoặc tiến trình con khi dev.",
         "POST /api/analyze/filter tính lại và kiểm chứng lại biểu đồ cho lọc chéo (tối đa 12 chart).",
         "Chỉ từ điển tổng hợp trong profiler.py nhìn vào tên cột."])
h("4.3 DB Agent", 2)
p("validator.py dùng sqlglot để chỉ cho phép SELECT và tự thêm LIMIT 1000; rls_transformer.py chèn điều kiện department_id/tenant_id. Khi đặt DB_AGENT_PASSWORD, gateway tạo vai trò Postgres "
  "chỉ đọc (SELECT trên DB_AGENT_TABLES, giao dịch read-only, statement timeout 15 s) và chính CSDL thực thi ranh giới ngay cả khi validator bị vượt qua. Truy vấn bảng nhạy cảm "
  "(lương/nhân sự/tài chính) bị chặn bởi HITL.")
h("4.4 Search Agent và Integration Agent", 2)
p("Search agent dùng Tavily (cần TAVILY_API_KEY) và Crawl4AI. Integration agent gọi dịch vụ ngoài qua MCP REST; mọi lời gọi có tác dụng ghi đều qua HITL.")
h("4.5 Human-in-the-Loop", 2)
p("Khi gặp thao tác nhạy cảm, executor lưu bản ghi chờ vào Orchestrator.pending_approvals (đường nhanh trong tiến trình) và phản chiếu vào Redis hitl:pending:<action_id> với TTL "
  "HITL_APPROVAL_TTL_SECONDS (900 s) để sống sót khi khởi động lại và dùng được trên mọi bản sao. Sự kiện human_approval_required được phát qua SSE; frontend gọi POST /api/chat/approve. "
  "_claim_pending đảm bảo đúng một bên xử lý; sau đó hệ thống thực thi lại với RLS (SQL) hoặc gọi MCP REST (tích hợp).")
h("4.6 Bộ nhớ và hội thoại", 2)
bullets(["Lịch sử phiên: Redis, tối đa 5 lượt; khóa phiên = <user>:<session client>.",
         "Bộ nhớ dài hạn mem0: chạy trong luồng riêng với timeout thật (ghi 30 s, đọc 5 s), che PII (số điện thoại VN, CCCD, email, thẻ tín dụng) trước khi lưu.",
         "Hội thoại người dùng đăng nhập đồng bộ qua /api/conversations với khóa lạc quan theo updated_at (so sánh đến mili giây)."])
page_break()

# ---------------------------------------------------------------- 5
h("5. Kiến trúc cơ sở dữ liệu và từ điển dữ liệu")
p("PostgreSQL chứa ba nhóm bảng: kho tri thức (rag_chunks), hội thoại (conversations) và bộ nhớ dài hạn của mem0 (mem0_memories khi MEM0_VECTOR_STORE=pgvector). "
  "Lược đồ được quản lý bằng Alembic (migrations/versions: 0001 = baseline rag_chunks, 0002 = conversations); gateway tự áp dụng khi khởi động (AUTO_MIGRATE) và dùng advisory lock để các bản sao không đua nhau.")
figure(DIAG / "erd_diagram.png", "Sơ đồ thực thể – quan hệ")
h("5.1 Từ điển dữ liệu", 2)
table(["Bảng", "Cột", "Kiểu", "Ràng buộc", "Mô tả"], [
    ["rag_chunks", "id", "khóa tự sinh", "PK", "Định danh chunk."],
    ["rag_chunks", "content", "TEXT", "NOT NULL", "Nội dung chunk (có thể kèm [Context: …])."],
    ["rag_chunks", "raw_content", "TEXT", "NOT NULL", "Nội dung gốc, dùng cho trích dẫn."],
    ["rag_chunks", "embedding", "VECTOR(EMBEDDING_DIM)", "HNSW cosine (m=16, ef_construction=64)", "Vector nhúng."],
    ["rag_chunks", "tsv", "tsvector", "GENERATED ALWAYS … STORED; GIN", "Full-text từ content (cấu hình english)."],
    ["rag_chunks", "filename, category, section_title, detected_pattern", "TEXT", "category có chỉ mục", "Siêu dữ liệu tài liệu và mẫu phát hiện."],
    ["rag_chunks", "doc_key", "TEXT", "chỉ mục", "'<category>/<filename>' – danh tính tài liệu."],
    ["rag_chunks", "doc_hash", "TEXT", "", "sha256 văn bản trích xuất, phát hiện thay đổi."],
    ["rag_chunks", "chunk_index, page", "INT", "page NULL với DOCX", "Thứ tự chunk và trang PDF bắt đầu."],
    ["rag_chunks", "tenant_id", "TEXT", "NOT NULL DEFAULT 'public'", "Phạm vi tenant cho RLS."],
    ["rag_chunks", "created_at", "TIMESTAMPTZ", "DEFAULT NOW()", "Thời điểm tạo."],
    ["conversations", "tenant_id, user_id, id", "TEXT", "PK ghép, NOT NULL", "Chủ sở hữu và id hội thoại."],
    ["conversations", "title", "TEXT", "NOT NULL DEFAULT ''", "Tiêu đề."],
    ["conversations", "pinned", "BOOLEAN", "NOT NULL DEFAULT FALSE", "Ghim."],
    ["conversations", "messages", "JSONB", "NOT NULL DEFAULT '[]'", "Danh sách tin nhắn."],
    ["conversations", "created_at, updated_at", "TIMESTAMPTZ", "NOT NULL DEFAULT NOW()", "updated_at dùng cho khóa lạc quan."],
], [1.0, 1.4, 1.3, 1.5, 1.6], 8.5)
h("5.2 Chỉ mục", 2)
bullets(["idx_rag_chunks_hnsw (HNSW, vector_cosine_ops) thay chỉ mục ivfflat trùng lặp đã bị xóa.",
         "idx_rag_chunks_tsv (GIN) làm full-text có chỉ mục; idx_rag_chunks_doc_key, idx_rag_chunks_category.",
         "idx_conversations_owner (tenant_id, user_id, updated_at DESC) cho danh sách hội thoại mới nhất.",
         "Vai trò db_agent có chính sách RLS tenant_scope; mỗi truy vấn chạy trong giao dịch đặt app.tenant_id/app.department_id."])
callout("Cảnh báo", "src/ingestion/schema.py đã ĐÓNG BĂNG (baseline 0001). Thay đổi lược đồ mới phải qua alembic revision -m \"...\".")
page_break()

# ---------------------------------------------------------------- 6
h("6. Đặc tả API và tích hợp")
p("Mọi route /api/* phụ thuộc authenticate (AUTH_MODE=off: người dùng ẩn danh cho dev; jwt: Bearer HS256 hoặc JWKS). Swagger có tại /docs. Lỗi trả theo ngôn ngữ của header X-UI-Lang.")
table(["Method", "URL", "Auth", "Mô tả"], [
    ["POST", "/api/chat", "Bearer + rate limit chat", "Chat đồng bộ JSON."],
    ["POST", "/api/chat/stream", "Bearer + rate limit chat", "Chat SSE: pev_step, plan, executing, verifying, answer_delta, human_approval_required, final_response."],
    ["POST", "/api/chat/approve", "Bearer, vai trò HITL_APPROVER_ROLES", "Phê duyệt/từ chối thao tác nhạy cảm."],
    ["POST", "/api/chat/title", "Bearer + rate limit chat", "Sinh tiêu đề hội thoại."],
    ["POST", "/api/analyze", "Bearer + rate limit analyze", "Tải CSV, trả dashboard_spec v2."],
    ["POST", "/api/analyze/filter", "Bearer", "Tính lại biểu đồ/KPI cho dòng đã lọc."],
    ["POST", "/api/knowledge/upload", "Bearer, KNOWLEDGE_UPLOAD_ROLES", "Thêm/cập nhật tài liệu vào kho RAG."],
    ["GET", "/api/knowledge/documents", "Bearer", "Liệt kê tài liệu (doc_key, chunks, pages)."],
    ["DELETE", "/api/knowledge/documents?doc_key=", "Bearer, KNOWLEDGE_UPLOAD_ROLES", "Xóa tài liệu và bản sao."],
    ["GET", "/api/knowledge/page?doc_key&page&q", "Bearer", "Trang PDF dạng PNG có tô sáng đoạn trích dẫn (header X-Highlighted)."],
    ["GET/PUT/DELETE", "/api/conversations[/{id}]", "Bearer", "Danh sách, tạo/cập nhật (khóa lạc quan), xóa hội thoại."],
    ["POST", "/api/auth/login", "Giới hạn theo IP", "Đăng nhập, trả JWT."],
    ["GET", "/api/auth/me", "Bearer", "Thông tin người dùng hiện tại."],
    ["GET", "/health", "Không", "Liveness."],
    ["GET", "/ready", "Không", "Readiness: Postgres, Redis, reranker (503 khi Postgres/Redis/orchestrator hỏng)."],
], [0.9, 2.0, 1.6, 2.0], 8.5)
h("6.1 Ví dụ chi tiết", 2)
callout("Lưu ý", "Các payload dưới đây là ví dụ minh họa theo schema trong src/gateway/schemas.py; giá trị là dữ liệu giả.")
h("POST /api/chat", 3)
code('Request:\n{ "query": "Điều khoản bảo mật trong NDA kéo dài bao lâu?", "session_id": "s-001" }\n\n'
     '200 OK:\n{ "session_id": "s-001", "response": "Nghĩa vụ bảo mật kéo dài 5 năm kể từ ngày chấm dứt [1].",\n'
     '  "sources": [ { "file": "nda_template.pdf", "section": "Confidentiality", "category": "nda", "page": 3 } ] }\n\n'
     '400 Bad Request (prompt injection bị chặn):\n{ "detail": "Yêu cầu bị từ chối bởi bộ lọc an toàn." }\n\n'
     '429 Too Many Requests: header Retry-After; vượt RATE_LIMIT_CHAT_PER_MINUTE (mặc định 30).')
h("POST /api/chat/approve", 3)
code('Request:\n{ "session_id": "s-001", "action_id": "act-7f3a", "decision": "approve", "feedback": null }\n\n'
     '200 OK:\n{ "status": "completed", "action_id": "act-7f3a", "decision": "approve", "message": "..." }\n\n'
     '403 Forbidden: người gọi không có vai trò approver/admin.   404/409: action hết hạn hoặc đã được xử lý.')
h("POST /api/analyze (multipart/form-data)", 3)
code('Fields: file=<csv|xlsx|parquet|json>, session_id, language=vi|en\n\n'
     '200 OK:\n{ "session_id": "s-001", "explanation": "...", "dashboard_spec": { "version": 2, "story": [...], "facts": [...],\n'
     '  "kpis": [...], "charts": [ { "spec": {...}, "option": {...}, "insight": "..." } ], "slicers": [...], "table": {...} } }\n\n'
     '400: tệp rỗng/định dạng không hỗ trợ.   413: vượt giới hạn dung lượng (kiểm tra từ Content-Length).')
h("6.2 Tích hợp ngoài", 2)
table(["Dịch vụ", "Giao thức", "Mục đích"], [
    ["OpenRouter qua LiteLLM", "HTTPS (OPENROUTER_API_KEY)", "Gọi LLM nhanh/mạnh."],
    ["TEI reranker", "HTTP (RERANKER_ENDPOINT)", "Rerank cross-encoder."],
    ["Tavily / Crawl4AI", "HTTPS", "Tìm kiếm và thu thập web."],
    ["MCP REST", "HTTP", "Truy vấn DB và gọi tích hợp."],
    ["Langfuse", "HTTP", "Truy vết LLM (LANGFUSE_*)."],
    ["IdP (tùy chọn)", "JWKS", "Xác minh RS256/ES256."],
], [1.8, 2.0, 2.7])
page_break()

# ---------------------------------------------------------------- 7
h("7. Giao diện người dùng")
callout("Lưu ý", "Kho mã hiện không chứa ảnh chụp màn hình (docs/screenshots rỗng, không có tài nguyên ảnh nào ngoài thư viện bên thứ ba). "
        "Phần này mô tả giao diện theo mã nguồn frontend/; hãy thêm ảnh vào docs/screenshots rồi bổ sung vào tài liệu.")
table(["Màn hình/Thành phần", "Tệp", "Chức năng và hành trình"], [
    ["Đăng nhập", "app/login, app/api/auth/*", "Nhập tài khoản → JWT lưu cookie httpOnly → vào khung chat."],
    ["Chat", "ChatInterface, ChatInput, ChatMessage", "Hỏi đáp, chọn tác tử (AgentSelectorInChat), đính kèm CSV hoặc tài liệu, Markdown."],
    ["PEV Stepper", "PEVStepper, lib/pevTrace.ts", "Hiển thị planner/executor/verifier theo thời gian thực từ SSE; mô tả nút lấy từ từ điển i18n."],
    ["Nguồn trích dẫn", "SourcesList, PageViewer", "Danh sách nguồn; nút 'xem trang' mở trang PDF có tô sáng."],
    ["Phê duyệt", "ApprovalCard", "Thẻ Approve/Reject khi nhận human_approval_required."],
    ["Dashboard", "dashboard/DynamicDashboard, EChartComponent", "Lưới 12 cột: KPI, biểu đồ, story, slicer lọc chéo (gọi /api/analyze/filter); chỉ vẽ số liệu từ server."],
    ["Tóm tắt dữ liệu", "DataSummaryView", "SQL playground DuckDB-WASM trên CSV đã phân tích."],
    ["Kho tri thức", "KnowledgeDocumentsCard (/docs)", "Liệt kê, xóa tài liệu; tải lên bằng cách đính kèm trong chat."],
    ["Thanh bên", "Sidebar", "Danh sách hội thoại, đồng bộ qua /api/conversations (useConversationSync)."],
], [1.4, 2.0, 3.1], 9)
p("Toàn bộ chữ giao diện (vi/en) nằm ở lib/locales/vi.ts (hợp đồng khóa) và en.ts, đọc bằng t(lang, key, vars); tests/unit/i18nCoverage.test.ts bắt khóa thiếu, placeholder lệch hoặc chữ Việt viết cứng.")
h("7.1 Hành trình tiêu biểu", 2)
bullets(["Hỏi tài liệu: nhập câu hỏi → PEV Stepper chạy → câu trả lời hiện dần (answer_delta là bản xem trước) → final_response thay thế sau kiểm chứng → bấm nguồn để xem trang PDF.",
         "Phân tích CSV: kéo thả tệp → dashboard dựng sẵn → chọn slicer để lọc chéo → hỏi tiếp bằng ngôn ngữ tự nhiên hoặc viết SQL trong DuckDB.",
         "Truy vấn nhạy cảm: hỏi về lương → thẻ phê duyệt → approver bấm Approve → kết quả được thực thi với RLS."])
page_break()

# ---------------------------------------------------------------- 8
h("8. Bảo mật, hiệu năng và khả năng mở rộng")
h("8.1 Xác thực và phân quyền", 2)
bullets(["JWT: AUTH_MODE=jwt, Bearer HS256 (INTERNAL_JWT_SECRET) hoặc JWKS; /api/auth/login phát token (AUTH_TOKEN_TTL_MINUTES=480), mật khẩu băm scrypt (scripts.make_user).",
         "Principal xác định khóa phiên, tenant/phòng ban cho RLS (current_scope) và vai trò: HITL_APPROVER_ROLES (approver, admin), KNOWLEDGE_UPLOAD_ROLES (admin).",
         "Frontend giữ JWT trong cookie httpOnly; route handler mới cần xác thực phải tự thêm Bearer (lib/backendAuth.ts::authHeaders)."])
h("8.2 Phòng thủ đầu vào/đầu ra LLM", 2)
bullets(["security.py: chuẩn hóa NFKC, bỏ ký tự độ rộng 0, giải mã payload base64, khớp mẫu jailbreak; đầu vào người dùng bọc bằng nonce.",
         "Canary token: mọi phản hồi/sự kiện SSE đi ra được kiểm tra rò rỉ.",
         "Nguồn RAG bọc phong bì nonce; mã LLM viết chạy trong sandbox AST allow-list + container cô lập (HMAC, fail closed).",
         "Che PII trước khi lưu bộ nhớ dài hạn."])
h("8.3 Dữ liệu và mã hóa", 2)
bullets(["DB agent: sqlglot SELECT-only, LIMIT 1000, vai trò Postgres chỉ đọc, RLS theo tenant/phòng ban, statement timeout 15 s.",
         "Bí mật chỉ nằm trong .env (POSTGRES_PASSWORD, LANGFUSE_*; không mặc định, không có trong compose); cổng hạ tầng bind 127.0.0.1; image backend non-root.",
         "Mã hóa kênh truyền (TLS) nên được kết thúc ở reverse proxy/ingress khi triển khai production."])
h("8.4 Rate limiting, cache và đồng thời", 2)
table(["Cơ chế", "Cấu hình mặc định", "Ghi chú"], [
    ["Chat/stream/title", "30 / phút / người dùng hoặc IP", "RATE_LIMIT_CHAT_PER_MINUTE; Redis cửa sổ cố định dùng chung cho mọi bản sao; fail open."],
    ["Analyze/upload", "10 / phút", "RATE_LIMIT_ANALYZE_PER_MINUTE."],
    ["Đăng nhập", "10 / phút / IP", "Chống brute force."],
    ["Lịch sử phiên", "5 lượt", "Redis."],
    ["Reranker", "RERANKER_TIMEOUT 0.8 s", "Circuit-breaker, lùi về thứ tự RRF khi quá hạn."],
    ["HITL", "TTL 900 s", "Redis + _claim_pending chống xử lý đôi."],
], [1.6, 2.0, 2.9])
bullets(["Hiệu năng truy hồi: HNSW + GIN; ngân sách thời gian tổng cho rerank; mem0 chạy trong luồng với timeout.",
         "Khả năng mở rộng ngang: docker-compose.scale.yml chạy nhiều backend sau nginx; scripts.replica_check xác minh trạng thái dùng chung; advisory lock tuần tự hóa migration và setup.",
         "scripts.loadtest đo tải nhẹ (--users, --duration, --chat N)."])
page_break()

# ---------------------------------------------------------------- 9
h("9. Triển khai và vận hành")
h("9.1 Thiết lập cục bộ", 2)
code("python -m venv venv && venv\\Scripts\\activate\n"
     "pip install -r requirements-dev.txt -c requirements.lock\n"
     "copy .env.example .env      # điền OPENROUTER_API_KEY, POSTGRES_PASSWORD, ...\n"
     "uvicorn src.gateway.main:app --host 0.0.0.0 --port 8000 --reload\n"
     "python -m scripts.run_ingestion     # nạp ./dataset vào pgvector\n"
     "cd frontend && npm install && npm run dev   # cổng 3000")
h("9.2 Docker Compose", 2)
code("docker compose up -d --build\n"
     "docker exec agent_backend python -m scripts.run_ingestion\n"
     "# dev hot-reload: -f docker-compose.yml -f docker-compose.dev.yml up --build\n"
     "# nhiều bản sao: docker compose -p scaletest -f docker-compose.yml -f docker-compose.scale.yml up -d --scale backend=3 backend lb")
table(["Dịch vụ", "Image", "Cổng"], [
    ["postgres", "ankane/pgvector (POSTGRES_IMAGE)", "5432 (127.0.0.1)"], ["redis", "redis:8-alpine", "6379 (127.0.0.1)"],
    ["backend", "agent-backend", "8000"], ["frontend", "agent-frontend", "3001"],
    ["langfuse-web / worker", "langfuse/langfuse:2, langfuse-worker:2", "3005"],
    ["tei-reranker", "text-embeddings-inference:cpu-1.2", "8080"], ["python-sandbox", "agent-sandbox", "nội bộ"],
], [1.8, 3.0, 1.7])
h("9.3 Biến môi trường chính", 2)
table(["Biến", "Ý nghĩa"], [
    ["OPENROUTER_API_KEY / OPENROUTER_BASE_URL", "Khóa và endpoint LLM."],
    ["FAST_LLM_MODEL / HEAVY_LLM_MODEL", "Mô hình nhanh / mạnh."],
    ["POSTGRES_URL, POSTGRES_USER/PASSWORD/DB, REDIS_URL", "Kết nối dữ liệu."],
    ["AUTH_MODE, INTERNAL_JWT_SECRET, AUTH_USERS", "Xác thực."],
    ["DB_AGENT_PASSWORD, DB_AGENT_TABLES", "Vai trò chỉ đọc của db_agent."],
    ["RERANKER_ENDPOINT, RERANKER_TIMEOUT, RERANK_POOL_K, HYBRID_CANDIDATES_K", "Truy hồi và rerank."],
    ["SANDBOX_URL", "Sandbox phân tích (rỗng = tiến trình con, chỉ dev)."],
    ["CORS_ALLOW_ORIGINS", "Nguồn được phép."],
    ["LANGFUSE_ENABLED / PUBLIC_KEY / SECRET_KEY / INIT_USER_*", "Truy vết."],
    ["TAVILY_API_KEY", "Bật tìm kiếm web."],
    ["MEM0_VECTOR_STORE", "pgvector (compose) hoặc memory (test)."],
    ["INSTALL_BROWSER", "Build arg: false để bỏ Chromium."],
], [3.0, 3.5])
h("9.4 CI/CD, giám sát", 2)
bullets(["GitHub Actions (.github/workflows/ci.yml): cài phụ thuộc theo lock, ruff check, docker compose config, pytest với Postgres pgvector (RAG_TEST_DSN); job riêng nạp dữ liệu và chạy run_rag_eval làm cổng (--min-hit5 0.55 --min-mrr 0.40).",
         "Kiểm thử: pytest (nhiều tệp độc lập), vitest, Playwright; scripts/run_offline_eval.py cho LLM-as-a-Judge.",
         "Giám sát: GET /health (liveness), GET /ready (readiness), Langfuse :3005 truy vết từng yêu cầu."])
page_break()

# ---------------------------------------------------------------- 10
h("10. Lộ trình và hạn chế đã biết")
h("10.1 Hạn chế", 2)
table(["Hạn chế", "Tác động", "Hướng xử lý"], [
    ["Chunk tải từ chat bị cắt >600 ký tự (phần sau không lưu)", "Mất phần cuối đoạn dài", "Đang báo cho người dùng; cân nhắc tách thay vì cắt (cần chạy lại run_rag_eval)."],
    ["Phụ thuộc nhiều lần gọi LLM mỗi lượt", "Độ trễ và chi phí", "Phân tầng mô hình, bỏ judge khi kiểm toán tất định đạt."],
    ["Rate limit Redis fail open", "Không giới hạn khi Redis sập", "Chấp nhận có chủ đích; cảnh báo log."],
    ["Full-text dùng cấu hình english", "Tiếng Việt chủ yếu dựa vào vector", "Đánh giá cấu hình tìm kiếm tiếng Việt riêng."],
    ["Sandbox rỗng = tiến trình con", "Cô lập yếu hơn", "Luôn đặt SANDBOX_URL ở production."],
    ["Chưa có ảnh chụp màn hình trong repo", "Tài liệu UI thiếu hình", "Bổ sung vào docs/screenshots."],
], [2.3, 1.8, 2.4])
h("10.2 Lộ trình", 2)
bullets(["Mở rộng bộ đánh giá RAG thủ công tiếng Việt và theo dõi hit@k/MRR theo từng lần đổi cấu hình.",
         "Contextual retrieval (KB_LLM_CONTEXT) sau khi đo lợi ích trên kho thật.",
         "Thêm chỉ số vận hành (Prometheus/OpenTelemetry) bên cạnh Langfuse.",
         "Kiểm thử tải đa bản sao có chat thật định kỳ; bổ sung bảng điều khiển chi phí LLM.",
         "Xem thêm docs/DATA_AGENT_REVIEW_AND_ROADMAP.md và docs/PROJECT_REVIEW_AND_ROADMAP.md."])

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print("saved", OUT)
